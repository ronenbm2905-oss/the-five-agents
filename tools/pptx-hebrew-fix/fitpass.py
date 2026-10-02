# -*- coding: utf-8 -*-
"""Shrink only the paragraphs the Arial swap pushed onto an extra line.

Everything else keeps its size, so the deck stays visually identical except
where the wider face would have broken a line that used to fit.
"""
import math, sys
from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.util import Pt
EMU = 12700

class M:
    def __init__(s, p):
        f = TTFont(p); s.upem = f['head'].unitsPerEm
        s.cmap = f.getBestCmap(); s.hmtx = f['hmtx']
    def w(s, t, sz):
        return sum(s.hmtx[s.cmap.get(ord(c)) or s.cmap.get(0x20)][0] for c in t) * sz / s.upem

SRC = {'Heebo': (M('fonts/Heebo-Regular.ttf'), M('fonts/Heebo-Bold.ttf')),
       'Secular One': (M('fonts/SecularOne-Regular.ttf'),) * 2}
AR = (M('/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf'),
      M('/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf'))

base = Presentation('fixed-heebo.pptx')      # original metrics, for the target line count
tgt  = Presentation(sys.argv[1])             # the Arial deck we adjust in place
fixed = []

for si, (bsl, tsl) in enumerate(zip(base.slides, tgt.slides), 1):
    for bsh, tsh in zip(bsl.shapes, tsl.shapes):
        if not bsh.has_text_frame or bsh.width is None:
            continue
        tf = bsh.text_frame
        li = tf.margin_left  if tf.margin_left  is not None else 91440
        ri = tf.margin_right if tf.margin_right is not None else 91440
        aw = (bsh.width - li - ri) / EMU
        if aw <= 0:
            continue
        for bp, tp in zip(tf.paragraphs, tsh.text_frame.paragraphs):
            bruns = [r for r in bp.runs if r.text.strip()]
            truns = [r for r in tp.runs if r.text.strip()]
            if not bruns or len(bruns) != len(truns):
                continue
            def width(runs, arial):
                tot = 0.0
                for r in runs:
                    sz = r.font.size.pt if r.font.size else 18.0
                    fn = r.font.name or 'Heebo'
                    b = bool(r.font.bold)
                    if arial:
                        tot += AR[1 if (b or fn == 'Secular One') else 0].w(r.text, sz)
                    else:
                        tot += SRC.get(fn, SRC['Heebo'])[1 if b else 0].w(r.text, sz)
                return tot
            wo = width(bruns, False)
            if wo <= 0:
                continue
            want = math.ceil(wo / aw)              # lines the original took
            wn = width(truns, True)
            if math.ceil(wn / aw) <= want:
                continue
            # scale down just enough to land back on `want` lines, with 2% slack
            scale = (want * aw) / wn * 0.98
            for r in truns:
                cur = r.font.size.pt if r.font.size else 18.0
                new = max(8.0, math.floor(cur * scale))
                r.font.size = Pt(new)
            fixed.append((si, tsh.shape_id, truns[0].text[:38],
                          round(bruns[0].font.size.pt if bruns[0].font.size else 18.0),
                          round(truns[0].font.size.pt)))

tgt.save(sys.argv[1])
print("paragraphs resized: %d" % len(fixed))
for s, sid, t, a, b in fixed:
    print("  slide %-3d shape %-5s %2dpt -> %2dpt  | %s" % (s, sid, a, b, t))
